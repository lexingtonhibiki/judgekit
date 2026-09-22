# -*- coding: utf-8 -*-
"""DeepSeek 三方打分管线：A/B 独立评分 → C 仲裁 → 结果自动填回 数据审核_v3.xlsx。

用法：python training/archive/run_deepseek_score.py [--workers 4]
前置：.env 里有 DEEPSEEK_KEY；工作簿里已有「情感打分_138」sheet（build_review_xlsx.py 生成）。
断点续跑：已完成 (id, pass) 存 training/archive/score_cache.jsonl，重跑只补缺。
成本：138 条×3 次 ≈ ¥0.4 以内，约 8-12 分钟。
"""
import json
import re
import sys
import threading
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
WB = ROOT / "数据审核_v3.xlsx"
SHEET = "情感打分_138"
CACHE = ROOT / "training" / "archive" / "score_cache.jsonl"
KEY = next(l.split("=", 1)[1].strip().strip('"') for l in open(ROOT / ".env", encoding="utf-8") if l.startswith("DEEPSEEK_KEY="))
API = "https://api.deepseek.com/v1/chat/completions"
MODEL = "deepseek-chat"

RUBRIC = """你是电商评论标注员，给下面的用户评论打一个 0-10 的情感权重分（0=非常负面，10=非常正面，可带一位小数）。

打分规则：
- 只评用户对商品/服务本身的态度；物流/快递算服务的一部分但权重低。
- 纯客观事实陈述（刚签收还没用、只确认发货、"等用了再来评"）不算特别正面，上限 6 分。
- 有具体使用效果细节的好评给 7 分以上；细节越具体、越个人化，分越高。
- 反讽/阴阳怪气按字面背后的真实情绪给低分；夸张修辞（"绝了""哭死"）看上下文分辨真假。
- 疑似刷单/模板文也按文本表面情绪打分，不需要判断真假。

参考锚点（必须校准到这个尺度）：
[6分] 收到了，快递小哥够快的，刚拆开其他同事就拿走了小的两瓶，她们说试试好用就上京东买，我也第一次用，好用再评论 →（仅过签收关+转述他人推荐，本人未评效果）
[8分] 整体不错,房间大,干净.服务也挺好的.酒店离市中心,海边都挺近…就是不知为何有蚊子 →（整体正面+小抱怨，抱怨扣分但翻不了盘）
[1分] 掰碎了扔给流浪狗，狗闻了闻，没吃，跑了 →（强烈负面，细节凿实）
[5分] 已经收到了，还没安装，等装好了再说 →（纯事实陈述，无情绪）

只输出 JSON，不要任何其他文字：
{"score": 数字, "reason": "不超过40字的打分理由"}

评论："""

ARBITER = """你是标注仲裁员。两个独立标注员对同一条评论的打分如下，请综合裁定最终分。

评论原文：{text}
标注员A：{a} 分。理由：{ar}
标注员B：{b} 分。理由：{br}

裁定规则：
- 双方分差 ≤2：取平均值（保留一位小数），理由合并成一句。
- 双方分差 >2：说明有一方忽略了某个细节，你要自己重读原文独立判定，不要和稀泥折中。
- 情绪类别按最终分：≤3.5 负面；3.5-6.5 中性（含纯事实陈述）；>6.5 正面。
- 处理建议：文本残缺、疑似刷单、反讽难判、或双方分差 >4 → "需人工复核"；否则 "通过"。

只输出 JSON，不要任何其他文字：
{"final_score": 数字, "band": "负面|中性|正面", "reason": "综合理由不超过50字", "action": "通过|需人工复核"}"""


def call_deepseek(prompt, temperature):
    body = json.dumps({"model": MODEL, "messages": [{"role": "user", "content": prompt}],
                       "temperature": temperature, "max_tokens": 300}).encode("utf-8")
    last = None
    for attempt in range(4):
        try:
            req = urllib.request.Request(API, data=body, headers={
                "Content-Type": "application/json", "Authorization": f"Bearer {KEY}"})
            with urllib.request.urlopen(req, timeout=120) as resp:
                data = json.loads(resp.read().decode("utf-8"))
            content = data["choices"][0]["message"]["content"]
            m = re.search(r"\{.*\}", content, re.DOTALL)
            if m:
                return json.loads(m.group(0))
            last = ValueError("no json in: " + content[:80])
        except Exception as e:
            last = e
        time.sleep(3 * (attempt + 1))
    raise RuntimeError(f"deepseek failed after retries: {last}")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    workers = 4
    if "--workers" in sys.argv:
        workers = int(sys.argv[sys.argv.index("--workers") + 1])
    items = [json.loads(l) for l in open(ROOT / "training" / "real_pools" / "clean" / "real_sentiment.jsonl", encoding="utf-8") if l.strip()]
    done = set()
    cache_lock = threading.Lock()
    if CACHE.exists():
        for l in open(CACHE, encoding="utf-8"):
            if l.strip():
                o = json.loads(l)
                done.add((o["id"], o["pass"]))
    print(f"items={len(items)} cached={len(done)}")

    def ask(iid, pas, prompt, temp):
        if (iid, pas) in done:
            return None
        out = call_deepseek(prompt, temp)
        with cache_lock, open(CACHE, "a", encoding="utf-8") as f:
            f.write(json.dumps({"id": iid, "pass": pas, "payload": out}, ensure_ascii=False) + "\n")
        return out

    # A/B 两轮独立评分
    ab_jobs = [(r["id"], p, RUBRIC + r["text"], 0.7) for r in items for p in ("A", "B")]
    with ThreadPoolExecutor(max_workers=workers) as ex:
        list(ex.map(lambda j: ask(*j), ab_jobs))

    # C 仲裁
    cache = {}
    for l in open(CACHE, encoding="utf-8"):
        if l.strip():
            o = json.loads(l)
            cache[(o["id"], o["pass"])] = o["payload"]
    c_jobs = []
    for r in items:
        a, b = cache.get((r["id"], "A")), cache.get((r["id"], "B"))
        if a and b:
            # 用 replace 链而非 str.format：ARBITER 末尾含字面 JSON 示例 {"final_score":...}，
            # 直接 .format 会把这些裸花括号当字段名而抛 KeyError。replace 不动模板文本，渲染结果与 .format 一致。
            p = ARBITER
            for k, v in (("{text}", r["text"]), ("{a}", a.get("score")), ("{ar}", a.get("reason")),
                         ("{b}", b.get("score")), ("{br}", b.get("reason"))):
                p = p.replace(k, str(v))
            c_jobs.append((r["id"], "C", p, 0.2))
    with ThreadPoolExecutor(max_workers=workers) as ex:
        list(ex.map(lambda j: ask(*j), c_jobs))

    # 填回工作簿
    for l in open(CACHE, encoding="utf-8"):
        if l.strip():
            o = json.loads(l)
            cache[(o["id"], o["pass"])] = o["payload"]
    wb = load_workbook(WB)
    ws = wb[SHEET]
    row_by_id = {ws.cell(row=r, column=2).value: r for r in range(5, ws.max_row + 1)}
    filled = skipped = 0
    for r in items:
        rn = row_by_id.get(r["id"])
        a, b, c = cache.get((r["id"], "A")), cache.get((r["id"], "B")), cache.get((r["id"], "C"))
        if not (a and b and c) or rn is None:
            skipped += 1
            continue
        ws.cell(row=rn, column=5, value=c.get("final_score"))
        # band 由 final_score 确定性推导（≤3.5负面/3.5-6.5中性/>6.5正面），不采信模型自报 band
        fs = c.get("final_score")
        if isinstance(fs, (int, float)):
            band = "负面" if fs <= 3.5 else ("中性" if fs <= 6.5 else "正面")
        else:
            band = c.get("band", "")
        ws.cell(row=rn, column=6, value=band)
        ws.cell(row=rn, column=8, value=a.get("score"))
        ws.cell(row=rn, column=9, value=a.get("reason", "")[:40])
        ws.cell(row=rn, column=10, value=b.get("score"))
        ws.cell(row=rn, column=11, value=b.get("reason", "")[:40])
        ws.cell(row=rn, column=12, value=c.get("reason", "")[:50])
        ws.cell(row=rn, column=13, value="⚠人工" if c.get("action") == "需人工复核" else "✓通过")
        filled += 1
    wb.save(WB)
    print(f"FILLED={filled} SKIPPED={skipped} -> {WB.name} / {SHEET}")


if __name__ == "__main__":
    main()
