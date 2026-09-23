# -*- coding: utf-8 -*-
"""T17 120亲标gold重测spam：05垃圾_120全120行用户直标为最高gold。

输入：
  training/abc_out/数据审核_v4_full.xlsx（05垃圾_120：头行自找"id"/"我的最终"，
    我的最终=垃圾71/正常49零空即用户直标；FK60+JD60）
  training/abc_out/gold_frozen.jsonl（190行；仅取其中spam行gold做冲突比对，
    口径：120直标赢，冲突覆盖旧51口径）
任务：Task单例classify，labels=[垃圾,正常]，
  instruction=用户SCNv2 spam句（与T16同一对象复用，不重打字禁改字；
  criteria/labels/兜底与SCNv2-spam侧一致）。
供应商：只载typesafe（缺key即停线NEEDS_CONTEXT，不硬跑）。逐条run_task
  （fallback=True；兜底命中记ok=False、pred=None，不充正确；失败每条重试3次；
  断点续跑读已落盘id跳过）。handoff链不跑不碰。models.yaml槽位含义不碰。
输出：
  training/abc_out/gold_spam120.jsonl（120行：{id,text,gold,source}，
    gold=用户直标原样；FK60+JD60）
  training/abc_out/spam120_scores.jsonl（{id,task,source,gold,pred,conf,
    correct,latency_ms,cost,ok,error,provider}，gold=120直标口径）
  + stdout报告（提示词全文贴头+acc+Wilson95%CI分源+与旧口径冲突数/id清单）。
约120 calls。

用法：
  python training/eval_spam120.py
  python training/eval_spam120.py --limit 2   # 冒烟（只跑前2条）
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from judgekit.engine import Task  # noqa: E402
from training import eval_jev_v4 as jev  # noqa: E402
from training.eval_jev_v4 import (  # noqa: E402  同一对象，逐字零漂移
    SPAM_CRITERIA,
    SPAM_FALLBACK,
)

XLSX_PATH = ROOT / "training" / "abc_out" / "数据审核_v4_full.xlsx"
GOLD_PATH = ROOT / "training" / "abc_out" / "gold_frozen.jsonl"
GOLD120_PATH = ROOT / "training" / "abc_out" / "gold_spam120.jsonl"
OUT_PATH = ROOT / "training" / "abc_out" / "spam120_scores.jsonl"
SHEET_KEY = "05垃圾"
PROVIDER_NAME = "typesafe"
TRIES = 3

# 用户SCNv2 spam句（第二人称身份版；与jev.SCNV2_SPAM/T16同一对象，不重打字禁改字）。
SPAM_PROMPT = jev.SCNV2_SPAM


def build_spam_task() -> Task:
    """spam单例classify Task（labels垃圾/正常，instruction为SCNv2 spam句逐字整体）。"""
    return Task(name="jev-spam120-scnv2", primitive="classify",
                labels=["垃圾", "正常"],
                label_descriptions=dict(jev.SCNV2_SPAM_LABEL_DESC),
                criteria=SPAM_CRITERIA, instruction=SPAM_PROMPT,
                input_field="text", provider=PROVIDER_NAME,
                fallback_rules={k: list(v) for k, v in SPAM_FALLBACK.items()})


def find_header(ws) -> tuple[int, dict]:
    """自找头行：含“id”与“我的最终”的首行；返回(行号,{列名:列号})。"""
    for r in range(1, min(15, ws.max_row + 1)):
        cols = {ws.cell(r, c).value: c
                for c in range(1, ws.max_column + 1)
                if ws.cell(r, c).value}
        if "id" in cols and "我的最终" in cols:
            return r, cols
    raise SystemExit("05表头行未找到（含id/我的最终）")


def load_spam120_rows(xlsx: Path = XLSX_PATH) -> list[dict]:
    """读05垃圾_120：返回[{id,text,gold,source}]（gold=我的最终直标原样，sheet序）。"""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    name = next(s for s in wb.sheetnames if SHEET_KEY in s)
    ws = wb[name]
    hr, cols = find_header(ws)
    c_id, c_txt = cols["id"], cols["文本"]
    c_src, c_fin = cols["来源"], cols["我的最终"]
    rows: list[dict] = []
    for r in range(hr + 1, ws.max_row + 1):
        rid = ws.cell(r, c_id).value
        if not rid:
            continue
        fin = ws.cell(r, c_fin).value
        rows.append({"id": str(rid), "text": ws.cell(r, c_txt).value or "",
                     "gold": fin, "source": ws.cell(r, c_src).value or ""})
    wb.close()
    return rows


def load_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with open(path, encoding="utf-8") as f:
        return [json.loads(l) for l in f if l.strip()]


def compute_conflicts(new_rows: list[dict],
                      gold_path: Path = GOLD_PATH) -> tuple[list[str], dict]:
    """与gold_frozen.jsonl spam行gold比对：返回(冲突id清单, frozen_by_id)。

    口径：120直标赢——冲突行gold取new_rows值（不回写gold_frozen）。
    """
    frozen = {r["id"]: r for r in load_jsonl(gold_path)
              if r.get("task") == "spam"}
    return ([r["id"] for r in new_rows
             if r["id"] in frozen and frozen[r["id"]].get("gold") != r["gold"]],
            frozen)


def main() -> None:
    ap = argparse.ArgumentParser(description="T17 120亲标gold重测spam（只跑spam 120行）")
    ap.add_argument("--xlsx", default=str(XLSX_PATH))
    ap.add_argument("--gold", default=str(GOLD_PATH))
    ap.add_argument("--gold120", default=str(GOLD120_PATH))
    ap.add_argument("--out", default=str(OUT_PATH))
    ap.add_argument("--tries", type=int, default=TRIES)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    print(f"[spam120提示词全文] {SPAM_PROMPT}", flush=True)
    rows = load_spam120_rows(Path(a.xlsx))
    n_fk = sum(1 for r in rows if r["id"].startswith("fk_"))
    n_jd = sum(1 for r in rows if r["id"].startswith("jd_"))
    n_spam = sum(1 for r in rows if r["gold"] == "垃圾")
    n_ok = sum(1 for r in rows if r["gold"] == "正常")
    print(f"[spam120集合] 05垃圾_120 n={len(rows)}（FK{n_fk}+JD{n_jd}；"
          f"直标垃圾{n_spam}+正常{n_ok}；预期120=FK60+JD60）", flush=True)
    bad = [r["id"] for r in rows if r["gold"] not in ("垃圾", "正常")]
    if bad:
        raise SystemExit(f"直标非二值混入：{bad[:10]}（共{len(bad)}行）")

    # 落gold_spam120.jsonl（直标原样；{id,text,gold,source}）
    out_g = Path(a.gold120)
    with open(out_g, "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps({"id": r["id"], "text": r["text"],
                                "gold": r["gold"], "source": r["source"]},
                               ensure_ascii=False) + "\n")
    print(f"[spam120gold] {len(rows)}行 -> {out_g}", flush=True)

    # 冲突统计（旧口径gold_frozen spam行；120直标赢，不回写旧文件）
    conflicts, _ = compute_conflicts(rows, Path(a.gold))
    print(f"[spam120冲突] 与gold_frozen-spam行gold不一致{len(conflicts)}行"
          f"（口径：120直标赢，旧文件不回写）：{','.join(sorted(conflicts))}",
          flush=True)

    eval_rows = [dict(r, task="spam") for r in rows]  # task仅内存供judge_row用
    if a.limit:
        eval_rows = eval_rows[:a.limit]
    providers = jev.load_providers_strict()  # 缺key在此停线，零调用
    task = build_spam_task()
    tasks = {"spam": task}

    out = Path(a.out)
    done = {r["id"]: r for r in load_jsonl(out)}  # 断点续跑：已落盘id跳过
    n_skip = sum(1 for r in eval_rows if r["id"] in done)
    if n_skip:
        print(f"[spam120] 断点续跑：已落盘{n_skip}行跳过", flush=True)
    with open(out, "a", encoding="utf-8") as f:
        for i, row in enumerate(eval_rows):
            if row["id"] in done:
                continue
            rec = jev.judge_row(row, tasks, providers, tries=a.tries)
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
            f.flush()
            done[rec["id"]] = rec
            if (i + 1) % 20 == 0 or (i + 1) == len(eval_rows):
                print(f"[spam120] 进度 {len(done)}/{len(eval_rows)}", flush=True)
    recs = [done[r["id"]] for r in eval_rows if r["id"] in done]
    n = len(recs)
    k = sum(1 for r in recs if r.get("correct"))
    ok = sum(1 for r in recs if r.get("ok"))
    lo, hi = jev.wilson(k, n)
    print(f"[spam120榜单] n={n} acc={k}/{n}={k / n * 100:.1f}% "
          f"95%CI[{lo * 100:.1f}%,{hi * 100:.1f}%] Jev直判ok={ok}/{n}", flush=True)
    for src in ("JD刷单", "FakeReview"):
        rs = [r for r in recs if r.get("source") == src]
        if not rs:
            continue
        kk = sum(1 for r in rs if r.get("correct"))
        loo, hii = jev.wilson(kk, len(rs))
        lat = sum(r.get("latency_ms", 0) for r in rs) / len(rs)
        cny_k = sum(r.get("cost", 0.0) for r in rs) / len(rs) * 1000
        print(f"[spam120分源] {src} n={len(rs)} acc={kk}/{len(rs)}={kk / len(rs) * 100:.1f}% "
              f"95%CI[{loo * 100:.1f}%,{hii * 100:.1f}%] "
              f"lat均值={lat:.0f}ms ¥/千次={cny_k:.4f}", flush=True)


if __name__ == "__main__":
    main()
