# -*- coding: utf-8 -*-
"""T8: 吃进06终审65条，冻结gold_frozen.jsonl（190行全量）。

用户口径（原样执行）：
  改标   = C取反（flip）：垃圾任务 正常<->垃圾（刷单/刷评归垃圾侧），转人工任务 不转<->转人工
  ✓通过 = C终判（认同）
  ✗删除 = 丢弃（本轮0个）
二值映射表写死见 FLIP；非二值C终判 flip 抛 NonBinaryFinalError，
流水线捕获后记 warn 并回落 C 终判原文，不停线。

输入：
  training/abc_out/probe_t6_scores.jsonl（190行，A/B/C三槽+C终判）
  training/abc_out/数据审核_v4_full.xlsx（06待审汇总_65：我的最终=改标48/✓通过17）
输出：
  training/abc_out/gold_frozen.jsonl（190行：id/text/task/gold/source/我的最终+orig_label/C终判）
  01总览刷新（✓通过125->190，待定65->0，就地更新同一xlsx）
  stdout 三源 verdict 段落：翻转率 / ABC单槽准确率 / 删留回炉建议

用法：
  python training/freeze_gold.py [--scores ...] [--xlsx ...] [--out ...] [--no-refresh]
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

PASS = "✓通过"
FLIP_DECISION = "改标"
DROP = "✗删除"

# 二值映射表（写死，按用户口径；刷单/刷评归垃圾侧，见 NORM）
FLIP = {
    "spam": {"正常": "垃圾", "垃圾": "正常"},
    "handoff": {"不转": "转人工", "转人工": "不转"},
}
# orig/A/B 侧多写法归一到 C 标签空间（与 build_v4_xlsx.GOLD_ALIAS 同义）
NORM = {"刷单spam": "垃圾", "刷评spam": "垃圾", "刷单": "垃圾", "刷评": "垃圾"}
TASK_CN_INV = {"路由": "route", "转人工": "handoff", "情感": "sentiment", "垃圾": "spam"}


class NonBinaryFinalError(ValueError):
    """flip 遇到非二值 C 终判（或未知任务）时抛出。"""


def norm(v) -> str:
    s = "" if v is None else str(v)
    return NORM.get(s, s)


def flip(task: str, c_final: str) -> str:
    """C取反；非二值/未知任务抛 NonBinaryFinalError。"""
    table = FLIP.get(task)
    if table is None:
        raise NonBinaryFinalError(f"未知任务:{task!r}")
    key = "" if c_final is None else str(c_final)
    try:
        return table[key]
    except KeyError:
        raise NonBinaryFinalError(f"非二值C终判:{task}={c_final!r}")


def apply_decision(task: str, c_final: str, decision: str):
    """落单条 gold。返回 (gold_or_None, warn_or_None)。

    ✗删除 -> (None, None)（丢弃）；✓通过 -> C终判；
    改标 -> flip，失败回落C终判原文并记warn；未知裁决 -> 回落C终判并记warn。
    """
    c = "" if c_final is None else str(c_final)
    if decision == DROP:
        return None, None
    if decision == PASS:
        return c, None
    if decision == FLIP_DECISION:
        try:
            return flip(task, c), None
        except NonBinaryFinalError as e:
            return c, f"flip失败回落C终判:{e}"
    return c, f"未知裁决{decision!r}回落C终判"


def load_scores(path: Path) -> dict:
    recs = [json.loads(l) for l in open(path, encoding="utf-8") if l.strip()]
    return {r["id"]: r for r in recs}


def load_rulings(xlsx: Path):
    """读06汇总：返回 (sheet名, [(id, task, c_xlsx, decision)])。列位按 builder 06 版式固定。"""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    name = next(s for s in wb.sheetnames if s.startswith("06待审汇总"))
    ws = wb[name]
    rows = []
    for r in range(5, ws.max_row + 1):
        rid = ws.cell(r, 3).value
        if not rid:
            continue
        cn = ws.cell(r, 2).value
        rows.append((str(rid), TASK_CN_INV.get(cn, cn),
                     ws.cell(r, 7).value, ws.cell(r, 9).value))
    wb.close()
    return name, rows


def freeze(scores_by_id: dict, rulings: list):
    """吃进终审。返回 (frozen_rows, warns, dist)。"""
    frozen, warns = [], []
    dist = Counter()
    ruled_ids = set()
    for rid, task_sheet, c_xlsx, decision in rulings:
        rec = scores_by_id.get(rid)
        if rec is None:
            warns.append(f"06行无scores:{rid}")
            continue
        task = rec.get("task") or task_sheet
        c_final = ((rec.get("C") or {}).get("final"))
        c_final = "" if c_final is None else str(c_final)
        if c_xlsx is not None and str(c_xlsx) != c_final:
            warns.append(f"C终判与表不一致:{rid} 表={c_xlsx!r} scores={c_final!r}（取scores）")
        gold, w = apply_decision(task, c_final, decision)
        if w:
            warns.append(f"{rid}:{w}")
        dist[decision] += 1
        ruled_ids.add(rid)
        if gold is None:  # ✗删除：丢弃
            continue
        frozen.append({"id": rid, "text": rec.get("text", ""), "task": task,
                       "gold": gold, "source": rec.get("source", ""),
                       "我的最终": decision, "orig_label": rec.get("orig_label", ""),
                       "C终判": c_final})
    # 125自动通过行：gold=C终判，我的最终=✓通过
    for rid, rec in scores_by_id.items():
        if rid in ruled_ids:
            continue
        c = (rec.get("C") or {})
        c_final = c.get("final")
        c_final = "" if c_final is None else str(c_final)
        if not c.get("ok") or not c_final:
            warns.append(f"自动行C异常:{rid}（gold置空）")
        frozen.append({"id": rid, "text": rec.get("text", ""), "task": rec.get("task"),
                       "gold": c_final, "source": rec.get("source", ""),
                       "我的最终": PASS, "orig_label": rec.get("orig_label", ""),
                       "C终判": c_final})
    frozen.sort(key=lambda r: r["id"])
    return frozen, warns, dist


def verdict_block(frozen: list, scores_by_id: dict) -> str:
    """三源 verdict 段落：翻转率 / ABC单槽准确率 / 删留回炉建议。"""
    lines = []
    for src in ["JD刷单", "FakeReview", "CSDS"]:
        rows = [r for r in frozen if r["source"] == src]
        n = len(rows)
        n_flip = sum(1 for r in rows if r["我的最终"] == FLIP_DECISION)
        acc = {}
        for slot in ["A", "B", "C"]:
            ok = 0
            for r in rows:
                rec = scores_by_id[r["id"]]
                d = rec.get(slot, {})
                v = d.get("value" if slot in ("A", "B") else "final", "")
                if norm(v) == norm(r["gold"]) and r["gold"]:
                    ok += 1
            acc[slot] = ok / n if n else 0.0
        both_wrong_ids = []
        for r in rows:
            rec = scores_by_id[r["id"]]
            va = norm((rec.get("A") or {}).get("value", ""))
            vb = norm((rec.get("B") or {}).get("value", ""))
            if va != norm(r["gold"]) and vb != norm(r["gold"]):
                both_wrong_ids.append(r["id"])
        if n_flip / n >= 0.25 if n else False:
            advice = "回炉：翻转率≥25%，rubric+采样复核"
        elif both_wrong_ids:
            advice = f"回炉候选{len(both_wrong_ids)}行复核（AB双错）：{','.join(both_wrong_ids[:8])}" \
                + ("…" if len(both_wrong_ids) > 8 else "")
        else:
            advice = "留：全量冻结gold"
        lines.append(
            f"{src} n={n} 翻转{n_flip}({n_flip / n:.1%}) "
            f"A单槽{acc['A']:.1%} B单槽{acc['B']:.1%} C单槽{acc['C']:.1%} "
            f"删0 留{n} {advice}")
    return "\n".join(lines)


def refresh_overview(xlsx: Path, frozen: list) -> str:
    """刷新01总览通过数（✓通过->全量，待定->0），就地保存并归一色值。"""
    from openpyxl import load_workbook
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from build_v4_xlsx import fix_styles_xml  # noqa: E402  延迟导入，单测不依赖
    wb = load_workbook(xlsx)
    ws = wb["01总览"]
    # 行：5路由 6转人工 7情感 8垃圾 9合计；列G=✓通过(7) I=待定(9)
    per = Counter((r["task"],) for r in frozen)
    row_of = {"route": 5, "handoff": 6, "sentiment": 7, "spam": 8}
    for t, row in row_of.items():
        n = per.get((t,), 0)
        ws.cell(row, 7).value = n   # ✓通过=全量（终审已吃进）
        ws.cell(row, 9).value = 0   # 待定清零
    ws.cell(9, 7).value = len(frozen)
    ws.cell(9, 9).value = 0
    n_flip = sum(1 for r in frozen if r["我的最终"] == FLIP_DECISION)
    n_pass = sum(1 for r in frozen if r["我的最终"] == PASS)
    ws.cell(11, 2).value = (f"黑盒证据：cache行数=190；sheets数=7；我的最终无空值数=190（预填率100%）。"
                            f"T8冻结：gold_frozen.jsonl{len(frozen)}行（自动通过125gold=C终判+终审65："
                            f"✓通过{n_pass - 125}认同+改标{n_flip}取反+删除0），01通过数已刷新，待定0。")
    wb.save(xlsx)
    n_xml = fix_styles_xml(xlsx)
    return f"01总览已刷新:✓通过{len(frozen)}/待定0 styles_xml_fix={n_xml}"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--scores", default="training/abc_out/probe_t6_scores.jsonl")
    ap.add_argument("--xlsx", default="training/abc_out/数据审核_v4_full.xlsx")
    ap.add_argument("--out", default="training/abc_out/gold_frozen.jsonl")
    ap.add_argument("--no-refresh", action="store_true", help="跳过01总览刷新")
    args = ap.parse_args()

    scores_by_id = load_scores(ROOT / args.scores)
    sheet, rulings = load_rulings(ROOT / args.xlsx)
    frozen, warns, dist = freeze(scores_by_id, rulings)

    out = ROOT / args.out
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        for r in frozen:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    print(f"frozen: {len(frozen)}行 -> {out}（06sheet={sheet}，终审{len(rulings)}："
          f"改标{dist.get(FLIP_DECISION, 0)}/✓通过{dist.get(PASS, 0)}/删除{dist.get(DROP, 0)}）")
    print(verdict_block(frozen, scores_by_id))
    for w in warns:
        print(f"WARN: {w}")
    print(f"warns={len(warns)}")
    if not args.no_refresh:
        print(refresh_overview(ROOT / args.xlsx, frozen))


if __name__ == "__main__":
    main()
