# -*- coding: utf-8 -*-
"""T16 spam单任务用户版重测(亲审gold)：只跑spam 51行。

输入：
  training/abc_out/数据审核_v4_full.xlsx（06待审汇总_65：任务列=垃圾51+转人工14；
    垃圾即spam，FK22+JD29；列位按builder 06版式：B列任务/C列id/G列C终判/I列我的最终）
  training/abc_out/gold_frozen.jsonl（190行；gold字段为用户亲审口径：
    改标=C取反/✓通过=C认同；本脚本gold直接取gold字段，不重算）
任务：Task单例classify，labels=[垃圾,正常]，
  instruction=用户SCNv2 spam句逐字（第二人称身份版，与eval_jev_v4.SCNV2_SPAM同一对象，
  不重打字防漂移；criteria/labels/兜底与SCNv2-spam侧一致，保证与SCNv2-spam子集可比）。
供应商：只载typesafe（缺key即停线NEEDS_CONTEXT，不硬跑）。逐条run_task
  （fallback=True；兜底命中记ok=False、pred=None，不充正确；失败每条重试3次；
  断点续跑读已落盘id跳过）。handoff链不跑不碰。
输出：training/abc_out/spam_only_scores.jsonl（{id,task,source,gold,pred,conf,
  correct,latency_ms,cost,ok,error,provider}）+ stdout报告（提示词全文贴头+
  acc+Wilson95%CI分源+与SCNv2-spam子集同id对比+74.2全量基线注记）。
约51 calls。

用法：
  python training/eval_spam_only.py
  python training/eval_spam_only.py --limit 2   # 冒烟（只跑前2条）
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from judgekit.engine import Task  # noqa: E402
from training import eval_jev_v4 as jev  # noqa: E402
from training.eval_jev_v4 import (  # noqa: E402  同一对象，逐字零漂移
    SPAM_CRITERIA,
    SPAM_FALLBACK,
)

XLSX_PATH = ROOT / "training" / "abc_out" / "数据审核_v4_full.xlsx"
GOLD_PATH = ROOT / "training" / "abc_out" / "gold_frozen.jsonl"
OUT_PATH = ROOT / "training" / "abc_out" / "spam_only_scores.jsonl"
SCNV2_PATH = ROOT / "training" / "abc_out" / "jev_v4_scnv2_scores.jsonl"
PROVIDER_NAME = "typesafe"
TRIES = 3

# 用户SCNv2 spam句逐字（第二人称身份版；与jev.SCNV2_SPAM同一对象，不重打字）。
SPAM_PROMPT = jev.SCNV2_SPAM

TASK_CN_TO_EN = {"垃圾": "spam", "spam": "spam",
                 "转人工": "handoff", "handoff": "handoff"}


def build_spam_task() -> Task:
    """spam单例classify Task（labels垃圾/正常，instruction为SCNv2 spam句逐字整体）。"""
    return Task(name="jev-spam-only-scnv2", primitive="classify",
                labels=["垃圾", "正常"],
                label_descriptions=dict(jev.SCNV2_SPAM_LABEL_DESC),
                criteria=SPAM_CRITERIA, instruction=SPAM_PROMPT,
                input_field="text", provider=PROVIDER_NAME,
                fallback_rules={k: list(v) for k, v in SPAM_FALLBACK.items()})


def load_spam_ids(xlsx: Path = XLSX_PATH) -> list[str]:
    """读06汇总（sheet名含“汇总”）：任务列归一后==spam的行id（sheet序，保持稳定）。"""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    name = next(s for s in wb.sheetnames if "汇总" in s)
    ws = wb[name]
    ids: list[str] = []
    for r in range(5, ws.max_row + 1):
        rid = ws.cell(r, 3).value
        if not rid:
            continue
        cn = ws.cell(r, 2).value
        task = TASK_CN_TO_EN.get("" if cn is None else str(cn).strip(), str(cn))
        if task == "spam":
            ids.append(str(rid))
    wb.close()
    return ids


def load_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with open(path, encoding="utf-8") as f:
        return [json.loads(l) for l in f if l.strip()]


def load_spam_rows(gold_path: Path = GOLD_PATH,
                   ids: list[str] | None = None) -> list[dict]:
    """取gold_frozen中对应id行（gold字段即亲审口径，不重算；保持ids序）。"""
    if ids is None:
        ids = load_spam_ids()
    by_id = {r["id"]: r for r in load_jsonl(gold_path)}
    rows = []
    for i in ids:
        r = by_id.get(i)
        if r is None:
            raise SystemExit(f"gold缺行：{i}（{gold_path}）")
        rows.append(r)
    return rows


def main() -> None:
    ap = argparse.ArgumentParser(description="T16 spam单任务用户版重测（亲审gold，只跑spam 51行）")
    ap.add_argument("--xlsx", default=str(XLSX_PATH))
    ap.add_argument("--gold", default=str(GOLD_PATH))
    ap.add_argument("--out", default=str(OUT_PATH))
    ap.add_argument("--scnv2", default=str(SCNV2_PATH),
                    help="SCNv2全量190分（同id对比spam子集用）")
    ap.add_argument("--tries", type=int, default=TRIES)
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()

    print(f"[spam-only提示词全文] {SPAM_PROMPT}", flush=True)
    ids = load_spam_ids(Path(a.xlsx))
    n_fk = sum(1 for i in ids if i.startswith("fk_"))
    n_jd = sum(1 for i in ids if i.startswith("jd_"))
    print(f"[spam-only集合] 06汇总spam行 n={len(ids)}（FK{n_fk}+JD{n_jd}；预期51=FK22+JD29）",
          flush=True)
    rows = load_spam_rows(Path(a.gold), ids)
    bad_task = [r["id"] for r in rows if r.get("task") != "spam"]
    if bad_task:
        raise SystemExit(f"gold task非spam混入：{bad_task}")
    if a.limit:
        rows = rows[:a.limit]
    providers = jev.load_providers_strict()  # 缺key在此停线，零调用
    task = build_spam_task()
    tasks = {"spam": task}

    out = Path(a.out)
    done = {r["id"]: r for r in load_jsonl(out)}  # 断点续跑：已落盘id跳过
    n_skip = sum(1 for r in rows if r["id"] in done)
    if n_skip:
        print(f"[spam-only] 断点续跑：已落盘{n_skip}行跳过", flush=True)
    with open(out, "a", encoding="utf-8") as f:
        for i, row in enumerate(rows):
            if row["id"] in done:
                continue
            rec = jev.judge_row(row, tasks, providers, tries=a.tries)
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
            f.flush()
            done[rec["id"]] = rec
            if (i + 1) % 20 == 0 or (i + 1) == len(rows):
                print(f"[spam-only] 进度 {len(done)}/{len(rows)}", flush=True)
    recs = [done[r["id"]] for r in rows if r["id"] in done]
    n = len(recs)
    k = sum(1 for r in recs if r.get("correct"))
    ok = sum(1 for r in recs if r.get("ok"))
    lo, hi = jev.wilson(k, n)
    print(f"[spam-only榜单] n={n} acc={k}/{n}={k / n * 100:.1f}% "
          f"95%CI[{lo * 100:.1f}%,{hi * 100:.1f}%] Jev直判ok={ok}/{n}", flush=True)
    for src in ("JD刷单", "FakeReview"):
        rs = [r for r in recs if r.get("source") == src]
        if not rs:
            continue
        kk = sum(1 for r in rs if r.get("correct"))
        loo, hii = jev.wilson(kk, len(rs))
        lat = sum(r.get("latency_ms", 0) for r in rs) / len(rs)
        cny_k = sum(r.get("cost", 0.0) for r in rs) / len(rs) * 1000
        print(f"[spam-only分源] {src} n={len(rs)} acc={kk}/{len(rs)}={kk / len(rs) * 100:.1f}% "
              f"95%CI[{loo * 100:.1f}%,{hii * 100:.1f}%] "
              f"lat均值={lat:.0f}ms ¥/千次={cny_k:.4f}", flush=True)
    # 与SCNv2-spam子集同id对比（同提示词全量190中的51行；74.2为全量190基线，口径不同仅注记）
    scn_by_id = {r["id"]: r for r in load_jsonl(Path(a.scnv2))}
    miss = [i for i in [r["id"] for r in recs] if i not in scn_by_id]
    if miss:
        print(f"[spam-only对比] SCNv2分缺行{len(miss)}（{miss[:5]}），对比跳过", flush=True)
    else:
        kk = sum(1 for r in recs if scn_by_id[r["id"]].get("correct"))
        nn = len(recs)
        loo, hii = jev.wilson(kk, nn)
        print(f"[spam-only对比] SCNv2-spam子集同51id n={nn} "
              f"acc={kk}/{nn}={kk / nn * 100:.1f}% 95%CI[{loo * 100:.1f}%,{hii * 100:.1f}%] "
              f"Δ(本轮-子集)={k - kk:+d}行/{(k - kk) / nn * 100:+.1f}pt；"
              f"74.2%为全量190基线（T14 SCN）口径不同仅注记", flush=True)
        flips = [r["id"] for r in recs
                 if r.get("ok") and scn_by_id[r["id"]].get("ok")
                 and r.get("pred") != scn_by_id[r["id"]].get("pred")]
        if flips:
            dist = Counter()
            for i in flips:
                old = scn_by_id[i].get("pred")
                new = done[i].get("pred")
                dist[f"{old}→{new}"] += 1
            print(f"[spam-only对比] 同ok翻转{len(flips)}行"
                  f"（{dict(dist)}）：{','.join(flips[:20])}"
                  + ("…" if len(flips) > 20 else ""), flush=True)
        else:
            print("[spam-only对比] 同ok翻转0行（与SCNv2-spam子集预测完全一致）", flush=True)


if __name__ == "__main__":
    main()
