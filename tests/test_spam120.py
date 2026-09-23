# -*- coding: utf-8 -*-
"""T17离线测试：05垃圾_120集合120行+直标原样落gold+冲突覆盖规则。全离线，不联网不读key。"""
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from training import eval_spam120 as m  # noqa: E402
from training import eval_jev_v4 as jev  # noqa: E402
from training import eval_spam_only as t16  # noqa: E402

EXP_SPAM = "场景：购物评价区。你的身份是潜在买家。垃圾判定：对于你而言，纯粹情绪称赞（刷好评返现）/谩骂（恶意攻击）等不提供真实消费反馈的均为垃圾；只看text字段独立判定"

# 120直标 vs gold_frozen-spam行gold 冲突清单（口径：120直标赢；见T17报告）
EXP_CONFLICTS = sorted(["fk_0001", "fk_0005", "fk_0006", "fk_0007", "fk_0023",
                        "fk_0049", "fk_0053", "jd_0010", "jd_0011", "jd_0038",
                        "jd_0039", "jd_0046", "jd_0048", "jd_0060"])


_XLSX_CACHE: dict | None = None


def _xlsx_all() -> dict:
    """05表全量{id:{text,gold,source}}（单次加载，多测共用）。"""
    global _XLSX_CACHE
    if _XLSX_CACHE is None:
        from openpyxl import load_workbook
        wb = load_workbook(ROOT / "training" / "abc_out" / "数据审核_v4_full.xlsx",
                           read_only=True, data_only=True)
        ws = wb["05垃圾_120"]
        hr, cols = m.find_header(ws)
        out = {}
        for r in range(hr + 1, ws.max_row + 1):
            rid = ws.cell(r, cols["id"]).value
            if not rid:
                continue
            out[str(rid)] = {"text": ws.cell(r, cols["文本"]).value or "",
                             "gold": ws.cell(r, cols["我的最终"]).value,
                             "source": ws.cell(r, cols["来源"]).value or ""}
        wb.close()
        _XLSX_CACHE = out
    return _XLSX_CACHE


def _xlsx_cell_by_id(rid: str) -> dict:
    cell = _xlsx_all().get(rid)
    assert cell is not None, f"05表缺行:{rid}"
    return cell


def test_spam120_id_set_120_fk60_jd60():
    rows = m.load_spam120_rows()
    assert len(rows) == 120
    assert len({r["id"] for r in rows}) == 120  # 去重120
    c = Counter(r["id"].split("_")[0] for r in rows)
    assert c["fk"] == 60 and c["jd"] == 60
    g = Counter(r["gold"] for r in rows)
    assert g["垃圾"] == 71 and g["正常"] == 49  # 直标分布，零空
    assert not any(r["gold"] is None or str(r["gold"]).strip() == "" for r in rows)
    sc = Counter(r["source"] for r in rows)
    assert sc["FakeReview"] == 60 and sc["JD刷单"] == 60
    # 无handoff/他任务混入
    assert not any(r["id"].startswith("csds_") for r in rows)


def test_gold120_verbatim_from_xlsx():
    """直标原样：gold_spam120.jsonl每行id/text/gold/source与05表逐字一致。"""
    rows = m.load_spam120_rows()
    assert len(rows) == 120
    for r in rows:
        assert set(r) == {"id", "text", "gold", "source"}
        cell = _xlsx_cell_by_id(r["id"])
        assert r["gold"] == cell["gold"]  # 直标原样，不取反不认同
        assert r["text"] == cell["text"]
        assert r["source"] == cell["source"]


def test_prompt_verbatim_singleton_t16():
    assert m.SPAM_PROMPT == EXP_SPAM
    assert m.SPAM_PROMPT is jev.SCNV2_SPAM  # 与T16同一对象复用，禁改字
    assert t16.SPAM_PROMPT is jev.SCNV2_SPAM
    t = m.build_spam_task()
    assert t.primitive == "classify"
    assert t.labels == ["垃圾", "正常"]
    assert t.instruction == EXP_SPAM  # 整体逐字，非前置叠加
    assert t.provider == "typesafe"
    assert t.input_field == "text"


def test_conflict_override_rule_120_wins():
    """冲突覆盖：120直标赢；清单14行且落盘gold取直标值。"""
    rows = m.load_spam120_rows()
    conflicts, frozen = m.compute_conflicts(rows)
    assert sorted(conflicts) == EXP_CONFLICTS
    assert len(conflicts) == 14
    new_by_id = {r["id"]: r for r in rows}
    for i in conflicts:
        assert frozen[i]["task"] == "spam"
        assert frozen[i]["gold"] != new_by_id[i]["gold"]  # 确为不一致
        assert new_by_id[i]["gold"] in ("垃圾", "正常")
    # 非冲突行与旧口径一致
    for r in rows:
        if r["id"] not in set(conflicts):
            assert frozen[r["id"]]["gold"] == r["gold"]
    # 落盘文件（若已生成）同样原样
    p = ROOT / "training" / "abc_out" / "gold_spam120.jsonl"
    if p.exists():
        disk = [json.loads(l) for l in open(p, encoding="utf-8") if l.strip()]
        assert len(disk) == 120
        by_id = {r["id"]: r for r in disk}
        for i in EXP_CONFLICTS:
            assert by_id[i]["gold"] == new_by_id[i]["gold"]  # 120赢
